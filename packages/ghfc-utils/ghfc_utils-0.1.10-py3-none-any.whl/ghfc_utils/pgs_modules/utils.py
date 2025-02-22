import csv
import gzip

import docx
import pandas as pd


def write_dataframe_to_docx(df: pd.DataFrame, doc_file: str) -> None:
    """A function to write the content of the dataframe df
    to a new docx document and saving it to doc_file.

    Args:
        df (pandas dataframe): the dataframe to save into a WORD document
        doc_file (string): Name of the word document to create
    """
    doc = docx.Document()
    t = doc.add_table(df.shape[0] + 1, df.shape[1])
    for j in range(df.shape[-1]):
        t.cell(0, j).text = df.columns[j]
    for i in range(df.shape[0]):
        for j in range(df.shape[-1]):
            t.cell(i + 1, j).text = str(df.values[i, j])
    doc.save(doc_file)


def guess_separator(file: str) -> str:
    """Function to guess the column separator used in txt file (gzipped or not).
    This funciton manages only the ' ', '\t' and ',' separators.

    Args:
        file (string): path to the file we want to gues the seprator from

    Returns:
        string: column separator used in the file
    """
    sniffer = csv.Sniffer()
    if is_gz_file(file):
        with gzip.open(file, "rt") as f:
            return sniffer.sniff(f.readline(), delimiters=[" ", "\t", ","]).delimiter
    else:
        with open(file, newline="") as f:
            return sniffer.sniff(f.readline(), delimiters=[" ", "\t", ","]).delimiter
    return "\t"


def is_gz_file(filepath: str) -> bool:
    """Function to check if a file is gzipped without relying on the file extension.

    Args:
        filepath (string): path to the file to check

    Returns:
        boolean: True if the file is gzipped, False else
    """
    with open(filepath, "rb") as test_f:
        return test_f.read(2) == b"\x1f\x8b"
